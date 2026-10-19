import os
from typing import List, Dict, Any, Optional
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.config_loader import get_abs_path, get_logger
from core.db.database import get_connection, init_db

logger = get_logger("sheet_manager")

# Exact columns defined in the Project Blueprint (Phase 2):
# [ID | Platform | Video Name / Hook | Full Caption | Creator Handle | Original URL | Local MP4 Path | Music Used | Repost Status | Posted URL]
SHEET_COLUMNS = [
    "ID",
    "Platform",
    "Video Name / Hook",
    "Full Caption",
    "Creator Handle",
    "Original URL",
    "Local MP4 Path",
    "Music Used",
    "Repost Status",
    "Posted URL"
]


class MasterSheetManager:
    """
    Automated Master Sheet Engine (saved_videos_log.xlsx):
    Creates, styles, reads, and bidirectionally syncs the Excel sheet with SQLite.
    """

    def __init__(self, excel_path: Optional[str] = None):
        init_db()
        self.excel_path = excel_path or get_abs_path("master_sheet")
        self.ensure_sheet_exists()

    def ensure_sheet_exists(self) -> str:
        """Creates saved_videos_log.xlsx with styled headers if it does not exist yet."""
        if not os.path.exists(self.excel_path):
            self._write_styled_workbook([])
            logger.info(f"Created master sheet at {self.excel_path}")
        return self.excel_path

    def _write_styled_workbook(self, rows: List[Dict[str, Any]]) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Saved Videos Log"

        # Header styling
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        ws.append(SHEET_COLUMNS)
        for col_idx in range(1, len(SHEET_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        ws.row_dimensions[1].height = 26

        # Status colors
        status_fills = {
            "Published": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "Ready to Post": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
            "Downloaded": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
            "Scraped": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
        }

        for r_idx, r in enumerate(rows, start=2):
            row_vals = [
                r.get("ID") or r.get("id", r_idx - 1),
                r.get("Platform") or r.get("platform", "Instagram"),
                r.get("Video Name / Hook") or r.get("video_name_hook", ""),
                r.get("Full Caption") or r.get("ai_caption") or r.get("full_caption", ""),
                r.get("Creator Handle") or r.get("creator_handle", "@unknown"),
                r.get("Original URL") or r.get("original_url", ""),
                r.get("Local MP4 Path") or r.get("ready_mp4_path") or r.get("local_mp4_path", ""),
                r.get("Music Used") or r.get("music_used", "Original Audio"),
                r.get("Repost Status") or r.get("repost_status", "Scraped"),
                r.get("Posted URL") or r.get("posted_url", "")
            ]
            ws.append(row_vals)
            status_val = str(row_vals[8])
            for c_idx in range(1, len(SHEET_COLUMNS) + 1):
                c = ws.cell(row=r_idx, column=c_idx)
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=(c_idx in (3, 4)))
                if c_idx == 9 and status_val in status_fills:
                    c.fill = status_fills[status_val]

        # Set clean column widths
        col_widths = [8, 14, 34, 48, 20, 42, 42, 26, 18, 38]
        for idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        wb.save(self.excel_path)

    def sync_db_to_excel(self) -> str:
        """Exports all SQLite saved_videos records into saved_videos_log.xlsx."""
        conn = get_connection()
        cur = conn.cursor()
        cur.execute('''
            SELECT id, platform, video_name_hook, full_caption, ai_caption,
                   creator_handle, original_url, local_mp4_path, ready_mp4_path,
                   music_used, repost_status, posted_url
            FROM saved_videos
            ORDER BY id ASC
        ''')
        db_rows = [dict(r) for r in cur.fetchall()]
        conn.close()

        formatted_rows = []
        for r in db_rows:
            caption_display = r["ai_caption"] if r.get("ai_caption") else r.get("full_caption", "")
            mp4_display = r["ready_mp4_path"] if r.get("ready_mp4_path") else r.get("local_mp4_path", "")
            formatted_rows.append({
                "ID": r["id"],
                "Platform": r["platform"],
                "Video Name / Hook": r["video_name_hook"],
                "Full Caption": caption_display,
                "Creator Handle": r["creator_handle"],
                "Original URL": r["original_url"],
                "Local MP4 Path": mp4_display,
                "Music Used": r["music_used"],
                "Repost Status": r["repost_status"],
                "Posted URL": r["posted_url"]
            })

        self._write_styled_workbook(formatted_rows)
        return self.excel_path

    def sync_excel_to_db(self) -> int:
        """Reads any manual edits made in saved_videos_log.xlsx back into SQLite."""
        if not os.path.exists(self.excel_path):
            return 0
        try:
            df = pd.read_excel(self.excel_path).fillna("")
        except Exception as e:
            logger.warning(f"Could not read Excel sheet: {e}")
            return 0

        conn = get_connection()
        cur = conn.cursor()
        synced = 0
        for _, row in df.iterrows():
            orig_url = str(row.get("Original URL", "")).strip()
            if not orig_url:
                continue
            platform = str(row.get("Platform", "Instagram")).strip() or "Instagram"
            hook = str(row.get("Video Name / Hook", "Untitled Reel")).strip()
            caption = str(row.get("Full Caption", "")).strip()
            creator = str(row.get("Creator Handle", "@unknown")).strip()
            local_mp4 = str(row.get("Local MP4 Path", "")).strip()
            music = str(row.get("Music Used", "Original Audio")).strip()
            status = str(row.get("Repost Status", "Scraped")).strip()
            posted_url = str(row.get("Posted URL", "")).strip()

            cur.execute("SELECT id FROM saved_videos WHERE original_url = ?", (orig_url,))
            existing = cur.fetchone()
            if existing:
                cur.execute('''
                    UPDATE saved_videos
                    SET platform = ?, video_name_hook = ?, full_caption = ?,
                        creator_handle = ?, local_mp4_path = ?, music_used = ?,
                        repost_status = ?, posted_url = ?
                    WHERE id = ?
                ''', (platform, hook, caption, creator, local_mp4, music, status, posted_url, existing["id"]))
            else:
                cur.execute('''
                    INSERT INTO saved_videos
                    (platform, video_name_hook, full_caption, creator_handle, original_url, local_mp4_path, music_used, repost_status, posted_url)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (platform, hook, caption, creator, orig_url, local_mp4, music, status, posted_url))
            synced += 1

        conn.commit()
        conn.close()
        return synced

    def upsert_video_record(
        self,
        original_url: str,
        platform: str = "Instagram",
        video_name_hook: str = "Saved Reel",
        full_caption: str = "",
        ai_caption: str = "",
        creator_handle: str = "@unknown",
        local_mp4_path: str = "",
        ready_mp4_path: str = "",
        music_used: str = "Original Audio",
        duration_sec: float = 0.0,
        repost_status: str = "Downloaded",
        posted_url: str = ""
    ) -> Dict[str, Any]:
        """Inserts or updates a video row in SQLite and immediately regenerates saved_videos_log.xlsx."""
        conn = get_connection()
        cur = conn.cursor()
        cur.execute("SELECT * FROM saved_videos WHERE original_url = ?", (original_url,))
        existing = cur.fetchone()

        if existing:
            ex = dict(existing)
            new_platform = platform or ex["platform"]
            new_hook = video_name_hook if video_name_hook != "Saved Reel" else (ex["video_name_hook"] or video_name_hook)
            new_caption = full_caption or ex["full_caption"]
            new_ai_caption = ai_caption or ex["ai_caption"]
            new_creator = creator_handle if creator_handle != "@unknown" else (ex["creator_handle"] or creator_handle)
            new_local = local_mp4_path or ex["local_mp4_path"]
            new_ready = ready_mp4_path or ex["ready_mp4_path"]
            new_music = music_used if music_used != "Original Audio" else (ex["music_used"] or music_used)
            new_dur = duration_sec or ex["duration_sec"]
            new_status = repost_status or ex["repost_status"]
            new_posted = posted_url or ex["posted_url"]

            cur.execute('''
                UPDATE saved_videos
                SET platform = ?, video_name_hook = ?, full_caption = ?, ai_caption = ?,
                    creator_handle = ?, local_mp4_path = ?, ready_mp4_path = ?,
                    music_used = ?, duration_sec = ?, repost_status = ?, posted_url = ?
                WHERE id = ?
            ''', (
                new_platform, new_hook, new_caption, new_ai_caption,
                new_creator, new_local, new_ready,
                new_music, new_dur, new_status, new_posted, ex["id"]
            ))
            row_id = ex["id"]
        else:
            cur.execute('''
                INSERT INTO saved_videos (
                    platform, video_name_hook, full_caption, ai_caption,
                    creator_handle, original_url, local_mp4_path, ready_mp4_path,
                    music_used, duration_sec, repost_status, posted_url
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                platform, video_name_hook, full_caption, ai_caption,
                creator_handle, original_url, local_mp4_path, ready_mp4_path,
                music_used, duration_sec, repost_status, posted_url
            ))
            row_id = cur.lastrowid

        conn.commit()
        cur.execute("SELECT * FROM saved_videos WHERE id = ?", (row_id,))
        updated_row = dict(cur.fetchone())
        conn.close()

        self.sync_db_to_excel()
        return updated_row

    def update_status_by_id(
        self,
        video_id: int,
        repost_status: str,
        posted_url: Optional[str] = None,
        ready_mp4_path: Optional[str] = None,
        music_used: Optional[str] = None,
        ai_caption: Optional[str] = None
    ) -> Optional[Dict[str, Any]]:
        conn = get_connection()
        cur = conn.cursor()
        cur.execute("SELECT * FROM saved_videos WHERE id = ?", (video_id,))
        row = cur.fetchone()
        if not row:
            conn.close()
            return None

        ex = dict(row)
        from datetime import datetime
        pub_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S") if repost_status == "Published" else ex.get("published_at", "")
        cur.execute('''
            UPDATE saved_videos
            SET repost_status = ?,
                posted_url = ?,
                ready_mp4_path = ?,
                music_used = ?,
                ai_caption = ?,
                published_at = ?
            WHERE id = ?
        ''', (
            repost_status,
            posted_url if posted_url is not None else ex["posted_url"],
            ready_mp4_path if ready_mp4_path is not None else ex["ready_mp4_path"],
            music_used if music_used is not None else ex["music_used"],
            ai_caption if ai_caption is not None else ex["ai_caption"],
            pub_at,
            video_id
        ))
        conn.commit()
        cur.execute("SELECT * FROM saved_videos WHERE id = ?", (video_id,))
        updated = dict(cur.fetchone())
        conn.close()
        self.sync_db_to_excel()
        return updated

    def get_all_rows(self, platform: Optional[str] = None, status: Optional[str] = None) -> List[Dict[str, Any]]:
        conn = get_connection()
        cur = conn.cursor()
        query = "SELECT * FROM saved_videos WHERE 1=1"
        params = []
        if platform:
            query += " AND LOWER(platform) = LOWER(?)"
            params.append(platform)
        if status:
            query += " AND LOWER(repost_status) = LOWER(?)"
            params.append(status)
        query += " ORDER BY id DESC"
        cur.execute(query, params)
        rows = [dict(r) for r in cur.fetchall()]
        conn.close()
        return rows
